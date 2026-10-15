"""Gestione del trading journal in Excel (openpyxl)."""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.market_config import session_label
from src.market_structure import is_countertrend

JOURNAL_PATH = Path("journal/trading_journal.xlsx")
ERROR_CATEGORIES = ("Esecuzione", "Setup", "Psicologia/Rischio")
HEADERS = ["ID", "Aperto", "Chiuso", "Simbolo", "Lato", "Entry", "SL", "TP",
           "Exit", "R", "Esito", "Setup", "Errori", "Note", "Sessione", "MAE",
           "Trend H4", "Trend H1", "Controtendenza"]


def init_journal(path: Path = JOURNAL_PATH) -> None:
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Trades"
    ws.append(HEADERS)
    wb.create_sheet("Metriche")
    wb.save(path)


def add_trade(symbol: str, side: str, entry: float, sl: float, tp: float,
              setup: str = "", note: str = "", trend: dict | None = None,
              path: Path = JOURNAL_PATH) -> int:
    """trend: output di market_structure.multi_tf_trend() al momento dell'ingresso."""
    init_journal(path)
    wb = load_workbook(path)
    ws = wb["Trades"]
    trade_id = ws.max_row  # header = riga 1
    ws.append([trade_id, datetime.now().isoformat(timespec="seconds"), None,
               symbol, side, entry, sl, tp, None, None, "OPEN", setup, "", note,
               session_label(datetime.utcnow()), None,
               trend and trend["H4"], trend and trend["H1"],
               (is_countertrend(side, trend["bias"]) if trend else None)])
    wb.save(path)
    return trade_id


def close_trade(trade_id: int, exit_price: float, errors: str = "",
                mae: float | None = None, path: Path = JOURNAL_PATH) -> float:
    """`errors`: vuoto oppure una categoria in ERROR_CATEGORIES (opz. `Categoria: dettaglio`)."""
    if errors and errors.split(":")[0].strip() not in ERROR_CATEGORIES:
        raise ValueError(f"Categoria errore non valida: usa una tra {ERROR_CATEGORIES}")
    wb = load_workbook(path)
    ws = wb["Trades"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == trade_id:
            side, entry, sl = row[4].value, row[5].value, row[6].value
            risk = abs(entry - sl)
            pnl = (exit_price - entry) if side.lower() == "buy" else (entry - exit_price)
            r = pnl / risk if risk else 0.0
            row[2].value = datetime.now().isoformat(timespec="seconds")
            row[8].value, row[9].value = exit_price, round(r, 2)
            row[10].value = "WIN" if r > 0 else "LOSS" if r < 0 else "BE"
            row[12].value = errors
            row[15].value = mae  # peggior escursione avversa (in prezzo)
            wb.save(path)
            update_metrics(path)
            return r
    raise ValueError(f"Trade {trade_id} non trovato")


def update_metrics(path: Path = JOURNAL_PATH) -> dict:
    wb = load_workbook(path)
    rs = [r[9].value for r in wb["Trades"].iter_rows(min_row=2) if r[9].value is not None]
    wins = [r for r in rs if r > 0]
    losses = [r for r in rs if r < 0]
    m = {
        "Trade chiusi": len(rs),
        "Win rate": len(wins) / len(rs) if rs else 0,
        "Expectancy (R)": sum(rs) / len(rs) if rs else 0,
        "Profit factor": sum(wins) / abs(sum(losses)) if losses else None,
        "R totale": sum(rs),
    }
    ws = wb["Metriche"]
    ws.delete_rows(1, ws.max_row)
    for k, v in m.items():
        ws.append([k, v])
    wb.save(path)
    return m
